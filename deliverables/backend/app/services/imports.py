"""
Bulk import: parse, map, validate in dry run, then commit (WBS 3.1, 4.5, 4.6,
behind the import wizard at 10.11).

Import is migration and backfill, not the intake path -- direct UI entry is
primary (see project memory opptrack-intake-not-excel), and every value this
module writes is marked with the lowest trust level in app/db/models/field_value.py,
below every connector and far below a human edit. That is the connector
precedence order the product document sets out, expressed as data rather than as
a convention.

The dry run is the point of the wizard. Nothing is created until a person has
seen what would fail: a row missing Customer or Disty ASP cannot have revenue
computed at all and is blocking (app/calc/severity.py); a value that will not
coerce or canonicalise is reported against the field and the raw string, never
repaired by guessing. Both the file and its findings are sealed into a snapshot
first, so the commit re-reads the bytes that were validated rather than trusting
a second upload to be identical.
"""

import csv
import io
from dataclasses import dataclass, field

from sqlalchemy.orm import Session

from app.calc.severity import ADVISORY, BLOCKING, classify_severity
from app.db.models.field_value import TRUST_LEVELS
from app.db.models.field_value import FieldValue as FieldValueRow
from app.db.models.finding import Finding
from app.db.models.opportunity import Opportunity
from app.services.intake import create_opportunity
from app.intake.canonicalize import canonicalize_row
from app.intake.coerce import coerce_date, coerce_numeric, coerce_percentage
from app.intake.provenance import wrap_row
from app.schemas.opportunity import OpportunityCreate
from app.services.state_transitions import seed_initial_state
from app.services.rubric_versions import is_forbidden

# Intake field -> the coercer its ParamSpec dtype implies (WBS 3.3).
NUMERIC_FIELDS = ("eau_kpcs", "unit_set", "disty_asp", "resale_asp")
DATE_FIELDS = ("mp_date",)
PERCENT_FIELDS = ("confidence",)

REQUIRED_FIELDS = tuple(
    name for name, f in OpportunityCreate.model_fields.items() if f.is_required()
)

# Header aliases seen in the real workbook and its field-map document. A header
# outside this table is not silently dropped: it is reported as unmapped, and the
# wizard lets a person map it by hand before the dry run is accepted.
HEADER_ALIASES: dict[str, str] = {
    "region": "region",
    "customer": "customer",
    "end customer": "end_customer", "end_customer": "end_customer", "oem": "end_customer",
    "project": "project", "project name": "project",
    "application": "application",
    "product line": "product_line", "product_line": "product_line",
    "part number": "part_number", "part#": "part_number", "part_number": "part_number", "part no": "part_number",
    "design status": "design_status", "design_status": "design_status", "status": "design_status",
    "stage": "stage",
    "m/p date": "mp_date", "mp date": "mp_date", "mp_date": "mp_date", "sop": "mp_date",
    "eau": "eau_kpcs", "eau (kpcs)": "eau_kpcs", "eau_kpcs": "eau_kpcs",
    "unit/set": "unit_set", "unit set": "unit_set", "unit_set": "unit_set",
    "disty asp": "disty_asp", "distributor asp": "disty_asp", "disty_asp": "disty_asp",
    "resale asp": "resale_asp", "resale_asp": "resale_asp",
    "confidence": "confidence", "confidence level": "confidence", "confidence_level": "confidence",
    "opportunity id": "external_id", "opportunity_id": "external_id", "external id": "external_id",
    "competitor part#": "competitor_part", "competitor part": "competitor_part", "competitor_part": "competitor_part",
    "owner": "owner",
    "comments": "evidence", "evidence": "evidence",
    "confidence rationale": "confidence_rationale", "confidence_rationale": "confidence_rationale",
}


@dataclass
class RowFinding:
    rule_id: str
    severity: str
    field: str | None
    raw_value: str | None
    message: str


@dataclass
class PreviewRow:
    index: int  # 1-based row number in the source file, header excluded
    values: dict
    findings: list[RowFinding] = field(default_factory=list)

    @property
    def blocking(self) -> bool:
        return any(f.severity == BLOCKING for f in self.findings)


@dataclass
class ParsedFile:
    headers: list[str]
    mapping: dict[str, str]          # header -> intake field
    unmapped_headers: list[str]
    rows: list[dict]                 # header -> raw cell


def parse_bytes(content: bytes, filename: str) -> ParsedFile:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        headers, raw_rows = _read_xlsx(content)
    else:
        headers, raw_rows = _read_csv(content)

    mapping: dict[str, str] = {}
    unmapped: list[str] = []
    for header in headers:
        target = HEADER_ALIASES.get(header.strip().lower())
        if target:
            mapping[header] = target
        else:
            unmapped.append(header)
    return ParsedFile(headers=headers, mapping=mapping, unmapped_headers=unmapped, rows=raw_rows)


def _read_csv(content: bytes) -> tuple[list[str], list[dict]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    headers = list(reader.fieldnames or [])
    return headers, [dict(row) for row in reader]


def _read_xlsx(content: bytes) -> tuple[list[str], list[dict]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # openpyxl is a dev extra, not a runtime dependency
        raise ValueError(
            "reading .xlsx needs openpyxl (pip install 'opptrack-backend[dev]'); "
            "export the sheet as CSV, or install it"
        ) from exc

    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(h) if h is not None else "" for h in next(rows, ())]
    out = []
    for values in rows:
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        out.append({headers[i]: values[i] for i in range(min(len(headers), len(values)))})
    return headers, out


def _coerce_field(target: str, raw) -> tuple[object, str | None]:
    """Returns (value, note). A note means the raw string did not coerce, and
    the value is None -- never a zero or a guessed date standing in for it."""
    if raw is None or str(raw).strip() == "":
        return None, None
    if target in NUMERIC_FIELDS:
        coerced = coerce_numeric(raw)
        return coerced.value, coerced.note or None
    if target in DATE_FIELDS:
        if hasattr(raw, "isoformat"):  # openpyxl hands back a real date
            return raw, None
        coerced = coerce_date(raw)
        return coerced.value, coerced.note or None
    if target in PERCENT_FIELDS:
        coerced = coerce_percentage(raw)
        return coerced.value, coerced.note or None
    return str(raw).strip(), None


def _identity_key(values: dict) -> tuple[str, str, str, str] | None:
    required = ("region", "customer", "project", "part_number")
    if any(not values.get(field) for field in required):
        return None
    return tuple(str(values[field]).strip().casefold() for field in required)


def validate_rows(
    parsed: ParsedFile,
    mapping: dict[str, str] | None = None,
    rubric_version=None,
    existing_opportunities: list | None = None,
) -> list[PreviewRow]:
    """The dry-run pass. Runs coercion (3.3), canonicalisation (3.4) and severity
    (4.2) over every row and reports, without writing anything."""
    mapping = mapping or parsed.mapping
    preview: list[PreviewRow] = []
    existing_by_external = {
        o.external_id.casefold(): o for o in (existing_opportunities or []) if o.external_id
    }
    existing_by_identity = {
        key: o for o in (existing_opportunities or []) if (key := _identity_key({
            "region": o.region, "customer": o.customer, "project": o.project, "part_number": o.part_number,
        })) is not None
    }
    seen_external: dict[str, int] = {}
    seen_identity: dict[tuple[str, str, str, str], int] = {}

    for i, raw_row in enumerate(parsed.rows, start=1):
        values: dict = {}
        findings: list[RowFinding] = []

        for header, target in mapping.items():
            value, note = _coerce_field(target, raw_row.get(header))
            values[target] = value
            if note:
                findings.append(RowFinding(
                    rule_id="COERCE", severity=ADVISORY, field=target,
                    raw_value=str(raw_row.get(header)), message=note,
                ))

        result = canonicalize_row({k: v for k, v in values.items() if v is not None})
        for finding in result.findings:
            findings.append(RowFinding(
                rule_id="CANON", severity=BLOCKING, field=finding.field,
                raw_value=str(finding.raw_value), message=finding.message,
            ))
        values.update(result.canonical)

        external_id = str(values.get("external_id") or "").strip()
        identity_key = _identity_key(values)
        duplicate_message = None
        if external_id and external_id.casefold() in existing_by_external:
            duplicate_message = f"external opportunity id {external_id!r} already maps to {existing_by_external[external_id.casefold()].id}"
        elif external_id and external_id.casefold() in seen_external:
            duplicate_message = f"external opportunity id {external_id!r} is duplicated by import row {seen_external[external_id.casefold()]}"
        elif identity_key in existing_by_identity:
            duplicate_message = f"business identity already maps to {existing_by_identity[identity_key].external_id}"
        elif identity_key in seen_identity:
            duplicate_message = f"business identity is duplicated by import row {seen_identity[identity_key]}"

        if duplicate_message:
            findings.append(RowFinding(
                rule_id="DUPLICATE",
                severity=BLOCKING,
                field="external_id" if external_id else "project",
                raw_value=external_id or values.get("project"),
                message=f"{duplicate_message}; review the crosswalk instead of creating a second opportunity",
            ))

        if external_id:
            seen_external.setdefault(external_id.casefold(), i)
        if identity_key is not None:
            seen_identity.setdefault(identity_key, i)

        if (
            rubric_version is not None
            and values.get("design_status") is not None
            and values.get("stage") is not None
            and is_forbidden(rubric_version, values["design_status"], values["stage"])
        ):
            findings.append(RowFinding(
                rule_id="MATRIX_A",
                severity=BLOCKING,
                field="design_status",
                raw_value=f"{values['design_status']} / {values['stage']}",
                message=(
                    f"design status {values['design_status']!r} cannot be paired with "
                    f"stage {values['stage']!r} under Matrix A version {rubric_version.label!r}"
                ),
            ))

        for required in REQUIRED_FIELDS:
            if values.get(required) is None:
                findings.append(RowFinding(
                    rule_id="REQUIRED", severity=BLOCKING, field=required, raw_value=None,
                    message=f"{required} is required and the row does not supply it",
                ))

        if classify_severity(values) == BLOCKING and not any(f.rule_id == "REQUIRED" for f in findings):
            findings.append(RowFinding(
                rule_id="SEVERITY", severity=BLOCKING, field=None, raw_value=None,
                message="row is missing a field revenue cannot be computed without (WBS 4.2)",
            ))

        preview.append(PreviewRow(index=i, values=values, findings=findings))

    # WBS 4.6 / 4.1: the same V-a..V-i rules the workbook report runs, over the
    # import rows, so a row that would be flagged on the sheet is flagged here
    # before it is written. Tab-scope findings (V-c, V-i) name every row they
    # cover, so each of those rows carries the finding.
    _apply_rules(preview, rubric_version)
    return preview


def _apply_rules(preview: list[PreviewRow], rubric_version) -> None:
    from app.guards.report import matrix_from_cells
    from app.guards.rules import RuleContext, run_rules
    from app.intake.xlsx_reader import Cell, RawRow

    raw_rows = []
    for row in preview:
        cells = {k: Cell(coordinate=f"import!row {row.index}", value=v, formula=None)
                 for k, v in row.values.items() if v is not None}
        raw_rows.append(RawRow(tab="import", row_number=row.index, fields=dict(row.values), cells=cells))
    allowed, baselines = (None, None)
    if rubric_version is not None:
        allowed, baselines = matrix_from_cells((rubric_version.matrix_a or {}).get("cells", {}))
    findings, _dark = run_rules(
        RuleContext(rows=raw_rows, matrix_a=allowed, baselines=baselines),
        rule_ids=("V-a", "V-c", "V-d", "V-g", "V-i"),
    )
    by_index = {row.index: row for row in preview}
    for f in findings:
        if f.rule_id == "V-a" and f.severity == BLOCKING:
            continue  # the REQUIRED/SEVERITY checks above already say this
        targets = [by_index[f.row_number]] if f.row_number in by_index else [
            r for r in preview if f.cell and f"row {r.index}" in f.cell
        ]
        for target in targets:
            if any(x.rule_id == f.rule_id and x.field == f.field and x.message == f.message for x in target.findings):
                continue
            target.findings.append(RowFinding(
                rule_id=f.rule_id, severity=f.severity if f.severity != "cosmetic" else ADVISORY,
                field=f.field, raw_value=f.raw_value, message=f.message,
            ))


def persist_findings(db: Session, snapshot_id: str, preview: list[PreviewRow]) -> int:
    """Findings are stored against the snapshot, which is what lets the wizard's
    verdict be re-read later rather than recomputed from a file someone may have
    edited since."""
    count = 0
    for row in preview:
        for finding in row.findings:
            db.add(Finding(
                opportunity_id=None, snapshot_id=snapshot_id,
                rule_id=finding.rule_id, severity=finding.severity,
                message=f"row {row.index}"
                        + (f", {finding.field}" if finding.field else "")
                        + f": {finding.message}",
            ))
            count += 1
    db.flush()
    return count


def commit_rows(db: Session, preview: list[PreviewRow], actor: str, snapshot_id: str | None = None) -> list[Opportunity]:
    """Creates an opportunity per non-blocking row, through the same
    canonicalise -> provenance -> seed-state path as the direct-entry endpoint.
    Blocking rows are not partially written: they are left out, and the caller
    reports how many and why."""
    created: list[Opportunity] = []
    for row in preview:
        if row.blocking:
            continue
        values = {k: v for k, v in row.values.items() if k in OpportunityCreate.model_fields}
        payload = OpportunityCreate(**values)
        # The same service direct entry uses (WBS 3.7); import is the lowest
        # trust tier, so a later human edit or connector value supersedes it.
        result = create_opportunity(
            db, payload.model_dump(), source="import", trust=TRUST_LEVELS["import"],
            actor=actor, snapshot_id=snapshot_id,
        )
        created.append(result.opportunity)

    db.flush()
    return created
