"""
Regression tests for the NRE line splitter (WBS 3.5).

test_real_fcst_revenue_splits_seven_nre_from_three_volume_rows reads
FCST_Revenue directly from the real workbook (same approach as
test_canonicalize_intake.py / test_coerce.py) rather than transcribing values by
hand, to avoid a transcription error passing as a "real data" proof.
"""

from pathlib import Path

import pytest

from app.calc.nre import reconciles_to_original, split_nre_from_volume

WORKBOOK = Path(__file__).resolve().parents[3] / "data" / "TrackF (1).xlsx"


def test_splits_nre_from_volume_by_design_status():
    rows = [
        {"project": "TK1", "part_number": "ID801", "design_status": "M/P",
         "quarterly_revenue_k": {"CY25": [0, 0, 138.45, 415.35]}},
        {"project": "ID803 2nd NRE", "part_number": "ID803", "design_status": "NRE",
         "quarterly_revenue_k": {"CY25": [0, 126, 0, 0]}},
    ]
    volume_rows, nre_charges = split_nre_from_volume(rows)

    assert len(volume_rows) == 1
    assert volume_rows[0].project == "TK1"
    assert len(nre_charges) == 1
    assert nre_charges[0].project == "ID803 2nd NRE"
    assert nre_charges[0].total_revenue_k == pytest.approx(126.0)


def test_nre_recognition_is_case_insensitive():
    rows = [{"project": "x", "part_number": "y", "design_status": "nre",
              "quarterly_revenue_k": {"CY25": [10, 0, 0, 0]}}]
    _, nre_charges = split_nre_from_volume(rows)
    assert len(nre_charges) == 1


def test_reconciles_to_original_true_when_split_is_complete():
    rows = [
        {"project": "a", "part_number": "p1", "design_status": "M/P", "quarterly_revenue_k": {"CY25": [100, 0, 0, 0]}},
        {"project": "b", "part_number": "p2", "design_status": "NRE", "quarterly_revenue_k": {"CY25": [50, 0, 0, 0]}},
    ]
    volume_rows, nre_charges = split_nre_from_volume(rows)
    assert reconciles_to_original(rows, volume_rows, nre_charges)


def test_reconciles_to_original_false_if_a_row_is_dropped():
    rows = [
        {"project": "a", "part_number": "p1", "design_status": "M/P", "quarterly_revenue_k": {"CY25": [100, 0, 0, 0]}},
        {"project": "b", "part_number": "p2", "design_status": "NRE", "quarterly_revenue_k": {"CY25": [50, 0, 0, 0]}},
    ]
    volume_rows, nre_charges = split_nre_from_volume(rows)
    assert not reconciles_to_original(rows, volume_rows, nre_charges[:0])  # simulate a dropped NRE row


@pytest.mark.skipif(not WORKBOOK.exists(), reason="TrackF (1).xlsx not present in this checkout")
def test_real_fcst_revenue_splits_seven_nre_from_three_volume_rows():
    import openpyxl

    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb["FCST_Revenue"]

    rows = []
    for r in range(5, 15):  # the ten real data rows, per the field map doc
        project = ws.cell(row=r, column=4).value
        part_number = ws.cell(row=r, column=7).value
        design_status = ws.cell(row=r, column=8).value
        cy25 = [ws.cell(row=r, column=c).value or 0 for c in (19, 20, 21, 22)]  # S:V
        cy26 = [ws.cell(row=r, column=c).value or 0 for c in (29, 30, 31, 32)]  # AC:AF
        rows.append({
            "project": project, "part_number": part_number, "design_status": design_status,
            "quarterly_revenue_k": {"CY25": cy25, "CY26": cy26},
        })

    assert len(rows) == 10  # "ten rows" per the field map doc

    volume_rows, nre_charges = split_nre_from_volume(rows)

    # "Seven of ten FCST_Revenue rows are non-recurring engineering charges" -- WBS 3.5's own count.
    assert len(nre_charges) == 7
    assert len(volume_rows) == 3

    assert reconciles_to_original(rows, volume_rows, nre_charges)

    # The two examples the field map doc cites by name and figure -- both are the
    # CY25 total (column W), the figure actually written in the doc; each also
    # carries a CY26 phase (168 and 210 respectively) the doc's prose doesn't
    # quote but this split still accounts for, which reconciles_to_original above
    # already proved ties to the workbook.
    by_project = {c.project: c for c in nre_charges}
    assert sum(by_project["ID803 2nd NRE"].quarterly_revenue_k["CY25"]) == pytest.approx(126.0)
    assert sum(by_project["Annual license(7 seats)"].quarterly_revenue_k["CY25"]) == pytest.approx(210.0)
