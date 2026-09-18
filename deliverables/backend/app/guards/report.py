"""
Findings report (WBS 4.5) -- the Phase 0 deliverable.

One function reads the workbook and returns every finding the rules produce;
one renders them as text. No database, no auth, no model: this is the data-
quality section of the specification generated from the file instead of
written by hand, and it is what `python -m scripts.findings_report` prints.

Matrix A comes from the published rubric when a database session is given,
otherwise from the v1 template (the decision register's grid), so the pairing
rule runs either way and says which table it used.
"""

from dataclasses import dataclass, field
from pathlib import Path

from app.guards.rules import (
    ADVISORY,
    BLOCKING,
    COSMETIC,
    RULES,
    TAB,
    Finding,
    context_from_workbook,
    run_rules,
)
from app.intake.xlsx_reader import Workbook, read_workbook
from app.services.rubric_versions import MATRIX_A_DEFAULTS, cell_key, matrix_a_template

SEVERITY_ORDER = {BLOCKING: 0, ADVISORY: 1, COSMETIC: 2}
SPEC_FINDINGS = tuple(f"F-{n:02d}" for n in range(1, 13))


@dataclass
class Report:
    workbook: str
    matrix_a_source: str
    findings: list[Finding] = field(default_factory=list)
    dark: dict[str, str] = field(default_factory=dict)
    row_counts: dict[str, int] = field(default_factory=dict)

    @property
    def by_severity(self) -> dict[str, int]:
        counts = {BLOCKING: 0, ADVISORY: 0, COSMETIC: 0}
        for f in self.findings:
            counts[f.severity] = counts.get(f.severity, 0) + 1
        return counts

    @property
    def spec_refs(self) -> set[str]:
        return {f.spec_ref for f in self.findings if f.spec_ref}

    @property
    def spec_missing(self) -> list[str]:
        return [ref for ref in SPEC_FINDINGS if ref not in self.spec_refs]

    def as_dict(self) -> dict:
        return {
            "workbook": self.workbook,
            "matrix_a_source": self.matrix_a_source,
            "row_counts": self.row_counts,
            "by_severity": self.by_severity,
            "spec_findings_reproduced": sorted(self.spec_refs),
            "spec_findings_missing": self.spec_missing,
            "dark_rules": self.dark,
            "findings": [
                {
                    "rule_id": f.rule_id, "severity": f.severity, "scope": f.scope, "where": f.where,
                    "tab": f.tab, "row": f.row_number, "field": f.field, "raw_value": f.raw_value,
                    "message": f.message, "spec_ref": f.spec_ref,
                }
                for f in self.findings
            ],
        }


def matrix_from_cells(cells: dict) -> tuple[dict[tuple[str, str], bool], dict[tuple[str, str], float]]:
    """A published (or template) Matrix A grid -> the two lookups the rules take."""
    allowed: dict[tuple[str, str], bool] = {}
    baselines: dict[tuple[str, str], float] = {}
    for key, cell in cells.items():
        status, stage = key.split("|", 1)
        if cell.get("allowed") is False:
            allowed[(status, stage)] = False
        elif cell.get("baseline") is not None:
            allowed[(status, stage)] = True
            baselines[(status, stage)] = float(cell["baseline"])
    return allowed, baselines


def _sheet1_findings(path: Path) -> list[Finding]:
    """F-11 lives on Sheet1, which the reader deliberately excludes (David Nam:
    a leftover from the template). Read it here, cosmetically, so the report
    still reproduces the documented finding."""
    import openpyxl

    try:
        ws = openpyxl.load_workbook(path, data_only=False)["Sheet1"]
    except KeyError:
        return []
    label_row = next((r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 4).value or "").strip().lower() == "deviation"), None)
    if label_row is None:
        return []
    empty = [ws.cell(label_row, c).coordinate for c in (5, 6) if ws.cell(label_row, c).value is None]
    computed = [ws.cell(label_row, c).coordinate for c in (7, 8, 9)
                if isinstance(ws.cell(label_row, c).value, str) and ws.cell(label_row, c).value.startswith("=")]
    if empty and computed:
        return [Finding(
            rule_id="V-h", severity=COSMETIC, scope=TAB, tab="Sheet1", row_number=label_row,
            cell=f"Sheet1!{', '.join(empty)}", field="deviation", raw_value=None,
            message=f"Deviation is computed for {', '.join(computed)} only; {', '.join(empty)} are empty. "
                    f"Sheet1 is a template leftover (decision #49/#50) and feeds nothing",
            spec_ref="F-11",
        )]
    return []


def workbook_report(path: str | Path, *, matrix_cells: dict | None = None, matrix_source: str | None = None) -> Report:
    """Reads the workbook and runs every rule. `matrix_cells` is a published
    Matrix A grid; without one the v1 template is used and labelled as such."""
    path = Path(path)
    wb: Workbook = read_workbook(path)
    if matrix_cells is None:
        matrix_cells = matrix_a_template()["cells"]
        matrix_source = matrix_source or "Matrix A v1 template (decision register, unpublished)"
    allowed, baselines = matrix_from_cells(matrix_cells)
    ctx = context_from_workbook(wb, matrix_a=allowed, baselines=baselines)
    findings, dark = run_rules(ctx)
    findings.extend(_sheet1_findings(path))
    findings.sort(key=lambda f: (SEVERITY_ORDER.get(f.severity, 9), f.rule_id, f.tab or "", f.row_number or 0))
    return Report(
        workbook=str(path), matrix_a_source=matrix_source or "published",
        findings=findings, dark=dark,
        row_counts={name: len(tab.rows) for name, tab in wb.tabs.items()},
    )


def render_text(report: Report) -> str:
    lines = [
        f"Findings report -- {report.workbook}",
        f"Matrix A: {report.matrix_a_source}",
        "Rows read: " + ", ".join(f"{k} {v}" for k, v in report.row_counts.items()),
        "",
        f"{report.by_severity[BLOCKING]} blocking, {report.by_severity[ADVISORY]} advisory, "
        f"{report.by_severity[COSMETIC]} cosmetic. Documented findings reproduced: "
        f"{len(report.spec_refs)} of {len(SPEC_FINDINGS)}"
        + (f" (missing {', '.join(report.spec_missing)})" if report.spec_missing else ""),
        "",
    ]
    width = max((len(f.where) for f in report.findings), default=10)
    for f in report.findings:
        ref = f" [{f.spec_ref}]" if f.spec_ref else ""
        lines.append(f"{f.severity.upper():9} {f.rule_id}  {f.where:<{width}}  {f.message}{ref}")
    if report.dark:
        lines.append("")
        lines.append("Rules that did not run: " + "; ".join(f"{k}: {v}" for k, v in report.dark.items()))
    lines.append("")
    lines.append("Rules: " + "; ".join(f"{r.id} {r.label}" for r in RULES))
    return "\n".join(lines)
